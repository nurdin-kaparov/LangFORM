"""Native coordinates resolve against original file bytes, never model values."""
from __future__ import annotations

import csv
import hashlib
import io
import json
from datetime import date, datetime, time
from pathlib import Path
from typing import Any


def column_id(column: int) -> str:
    result = ""
    while column:
        column, remainder = divmod(column - 1, 26)
        result = chr(65 + remainder) + result
    return result


def read_spreadsheet(path: str | Path) -> dict | None:
    path = Path(path)
    kind = path.suffix.lower().lstrip(".")
    if kind not in {"csv", "xlsx"}:
        return None
    data = path.read_bytes()
    if len(data) > 50 * 1024 * 1024:
        raise ValueError("Spreadsheet exceeds the 50 MB processing limit.")
    version = hashlib.sha256(data).hexdigest()
    if kind == "csv":
        rows = list(csv.reader(io.StringIO(data.decode("utf-8-sig"), newline=""), strict=True))
        grids = [{"sheet": None, "rows": rows, "columns": max(map(len, rows), default=0)}]
        if len(rows) * grids[0]["columns"] > 250_000:
            raise ValueError("Spreadsheet exceeds the 250,000 cell processing limit.")
    else:
        from openpyxl import load_workbook
        # data_only=False retains the source formula rather than stale cached results.
        workbook = load_workbook(io.BytesIO(data), data_only=False)
        if sum(sheet.max_row * sheet.max_column for sheet in workbook) > 250_000:
            workbook.close()
            raise ValueError("Spreadsheet exceeds the 250,000 cell processing limit.")
        grids = []
        for sheet in workbook:
            rows = []
            for row in sheet.iter_rows():
                values = []
                for cell in row:
                    merged = next((r for r in sheet.merged_cells.ranges if cell.coordinate in r
                                   and cell.coordinate != r.start_cell.coordinate), None)
                    value = {"mergedInto": merged.start_cell.coordinate} if merged else cell.value
                    if isinstance(value, (date, datetime, time)):
                        value = value.isoformat()
                    values.append(value)
                rows.append(values)
            # An untouched empty worksheet has no evidence.
            if not any(value is not None for row in rows for value in row):
                rows = []
            grids.append({"sheet": sheet.title, "rows": rows,
                          "columns": sheet.max_column if rows else 0})
        workbook.close()
    return {"type": kind, "version": version, "grids": grids}


def grid_coordinates(source: dict, grid: dict, start_row: int = 1, end_row: int | None = None,
                     start_column: int = 1, end_column: int | None = None) -> dict:
    end_row = len(grid["rows"]) if end_row is None else end_row
    end_column = grid["columns"] if end_column is None else end_column
    return {"type": source["type"], "sheet": grid["sheet"], "start_row": start_row,
            "end_row": end_row, "start_column": start_column, "end_column": end_column,
            "column_ids": [column_id(i) for i in range(start_column, end_column + 1)]}


def validate_coordinates(source: dict, value: Any) -> dict:
    if not isinstance(value, dict):
        raise ValueError("Invalid native coordinates.")
    grid = next((g for g in source["grids"] if g["sheet"] == value.get("sheet")), None)
    numbers = [value.get(key) for key in ("start_row", "end_row", "start_column", "end_column")]
    if (value.get("type") != source["type"] or grid is None or
        not all(type(n) is int for n in numbers)):
        raise ValueError("Invalid native coordinates.")
    sr, er, sc, ec = numbers
    if not 1 <= sr <= er <= len(grid["rows"]) or not 1 <= sc <= ec <= grid["columns"]:
        raise ValueError("Invalid native coordinates.")
    canonical = grid_coordinates(source, grid, sr, er, sc, ec)
    if "column_ids" in value and value["column_ids"] != canonical["column_ids"]:
        raise ValueError("Column identifiers do not match source coordinates.")
    return canonical


def contains(parent: dict, child: dict) -> bool:
    if parent["type"] != child["type"]:
        return False
    if parent["type"] == "paragraph":
        return parent["start_paragraph"] <= child["start_paragraph"] and parent["end_paragraph"] >= child["end_paragraph"]
    return (parent["sheet"] == child["sheet"] and parent["start_row"] <= child["start_row"]
            and parent["end_row"] >= child["end_row"] and parent["start_column"] <= child["start_column"]
            and parent["end_column"] >= child["end_column"])


def resolve_spreadsheet(source: dict, coordinates: dict, version: str | None) -> str:
    if not version or version != source["version"]:
        raise ValueError("Source changed after indexing. Analyze the Asset again.")
    c = validate_coordinates(source, coordinates)
    grid = next(g for g in source["grids"] if g["sheet"] == c["sheet"])
    cells = []
    for row in range(c["start_row"], c["end_row"] + 1):
        for column in range(c["start_column"], c["end_column"] + 1):
            values = grid["rows"][row - 1]
            value = values[column - 1] if column <= len(values) else None
            cells.append(f"{column_id(column)}{row} = {json.dumps(value, ensure_ascii=False)}")
    return "\n".join(cells)


def validate_grid_ranges(source: dict, ranges: Any, scope: dict | None = None) -> list[dict]:
    if scope:
        scope = validate_coordinates(source, scope)
    if not isinstance(ranges, list) or not ranges:
        raise ValueError("No native semantic ranges.")
    normalized = []
    by_key = {}
    for item in ranges:
        if (not isinstance(item, dict) or not isinstance(item.get("key"), str) or not item["key"]
            or item["key"] in by_key or not isinstance(item.get("reason"), str) or not item["reason"].strip()
            or "parentKey" not in item or item["parentKey"] is not None and not isinstance(item["parentKey"], str)):
            raise ValueError("Invalid or duplicate native range metadata.")
        c = validate_coordinates(source, item.get("coordinates"))
        if scope and not contains(scope, c):
            raise ValueError("Refinement lies outside parent coordinates.")
        entry = {**item, "coordinates": c}
        normalized.append(entry)
        by_key[entry["key"]] = entry

    def depth(item, seen):
        if item["key"] in seen:
            raise ValueError("Frame hierarchy contains a cycle.")
        if item.get("parentKey") is None:
            return 0
        parent = by_key.get(item["parentKey"])
        if not parent or not contains(parent["coordinates"], item["coordinates"]) or contains(item["coordinates"], parent["coordinates"]):
            raise ValueError("Child coordinates must lie strictly within parent.")
        return 1 + depth(parent, seen | {item["key"]})

    for item in normalized:
        item["depth"] = depth(item, set())
    for grid in source["grids"]:
        if not grid["rows"] or not grid["columns"] or scope and scope["sheet"] != grid["sheet"]:
            continue
        bounds = scope or grid_coordinates(source, grid)
        for row in range(bounds["start_row"], bounds["end_row"] + 1):
            intervals = sorted((item["coordinates"]["start_column"], item["coordinates"]["end_column"])
                               for item in normalized if item["coordinates"]["sheet"] == grid["sheet"]
                               and item["coordinates"]["start_row"] <= row <= item["coordinates"]["end_row"])
            next_column = bounds["start_column"]
            for start, end in intervals:
                if start > next_column:
                    break
                next_column = max(next_column, end + 1)
            if next_column <= bounds["end_column"]:
                raise ValueError("Native semantic ranges must cover the complete source.")
    return normalized
