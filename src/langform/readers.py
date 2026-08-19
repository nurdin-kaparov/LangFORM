from __future__ import annotations

import csv
import json
from pathlib import Path


class UnsupportedAssetTypeError(ValueError):
    pass


def read_text(path: str | Path) -> str:
    path = Path(path)
    suffix = path.suffix.lower()

    if suffix in {".txt", ".md", ".rst", ".py", ".log"}:
        return path.read_text(encoding="utf-8", errors="replace")

    if suffix == ".json":
        data = json.loads(path.read_text(encoding="utf-8"))
        return json.dumps(data, indent=2, ensure_ascii=False)

    if suffix == ".csv":
        rows = []
        with path.open("r", encoding="utf-8-sig", newline="", errors="replace") as f:
            reader = csv.reader(f)
            for row in reader:
                rows.append(" | ".join(row))
        return "\n".join(rows)

    if suffix == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as e:
            raise UnsupportedAssetTypeError(
                "PDF support requires: pip install -e \".[documents]\""
            ) from e
        reader = PdfReader(str(path))
        return "\n\n".join((page.extract_text() or "") for page in reader.pages)

    if suffix == ".docx":
        try:
            from docx import Document
        except ImportError as e:
            raise UnsupportedAssetTypeError(
                "DOCX support requires: pip install -e \".[documents]\""
            ) from e
        doc = Document(str(path))
        return "\n".join(p.text for p in doc.paragraphs)

    if suffix == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as e:
            raise UnsupportedAssetTypeError(
                "XLSX support requires: pip install -e \".[documents]\""
            ) from e
        wb = load_workbook(str(path), read_only=True, data_only=True)
        lines = []
        for ws in wb.worksheets:
            lines.append(f"# Sheet: {ws.title}")
            for row in ws.iter_rows(values_only=True):
                lines.append(" | ".join("" if v is None else str(v) for v in row))
        return "\n".join(lines)

    raise UnsupportedAssetTypeError(f"Unsupported file type: {suffix or '(no extension)'}")
