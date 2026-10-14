import json
from pathlib import Path

import pytest
from openpyxl import Workbook

from langform import ContextSubagent, ContextFrame, ContextPackage
from langform.coordinates import (
    read_spreadsheet, grid_coordinates, resolve_spreadsheet, validate_grid_ranges,
)
from langform.storage import JSONStore


@pytest.fixture(params=["csv", "xlsx"])
def spreadsheet(tmp_path, request):
    path = tmp_path / f"evidence.{request.param}"
    if request.param == "csv":
        path.write_bytes(Path("tests/fixtures/evidence.csv").read_bytes())
    else:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Sales '26"
        sheet.append(["Name", "Value", "Value", "Notes"])
        sheet.append(["alpha", "a,b", 10, "first\nsecond"])
        sheet.append(["beta", 'quote "inside"', 0, None])
        sheet.append([None, None, None, None])
        sheet.append(["gamma", None, 20, "=C2+C3"])
        sheet.merge_cells("A7:B7")
        sheet["A7"] = "merged"
        workbook.create_sheet("Other")["B3"] = "different sheet"
        workbook.save(path)
    return path


def native_ranges(source, scope=None):
    result = []
    for i, grid in enumerate(source["grids"]):
        if scope and scope["sheet"] != grid["sheet"]:
            continue
        parent = scope or grid_coordinates(source, grid)
        result.append({"key": f"parent{i}", "parentKey": None, "coordinates": parent, "reason": "Full useful table"})
        result.append({"key": f"child{i}", "parentKey": f"parent{i}",
                       "coordinates": grid_coordinates(source, grid, parent["start_row"] + 1,
                                                       parent["end_row"], parent["start_column"] + 1,
                                                       parent["end_column"]), "reason": "Value columns"})
    return result


def indexed(spreadsheet, tmp_path):
    service = ContextSubagent(tmp_path / "workspace")
    asset = service.register_asset(str(spreadsheet))
    service.analyze_asset(asset.asset_id)
    return service, asset


def test_original_cells_keep_positions_and_values(spreadsheet):
    source = read_spreadsheet(spreadsheet)
    grid = source["grids"][0]
    c = grid_coordinates(source, grid, 2, 3, 2, 4)
    content = resolve_spreadsheet(source, c, source["version"])
    assert 'B2 = "a,b"' in content
    assert 'B3 = "quote \\"inside\\""' in content
    assert 'C3 = 0' in content if source["type"] == "xlsx" else 'C3 = "0"' in content
    assert c["column_ids"] == ["B", "C", "D"]
    if source["type"] == "csv":
        assert len(grid["rows"]) == 5
    else:
        assert "=C2+C3" in resolve_spreadsheet(source, grid_coordinates(source, grid), source["version"])
        assert 'B7 = {"mergedInto": "A7"}' in resolve_spreadsheet(source, grid_coordinates(source, grid), source["version"])


def test_semantic_nested_ranges_not_row_chunks(spreadsheet, tmp_path):
    source = read_spreadsheet(spreadsheet)
    service = ContextSubagent(tmp_path / "workspace")
    frames = service.frame_builder.build_spreadsheet("asset", source, str(spreadsheet), planner=native_ranges)
    assert len(frames) == len(source["grids"]) * 2
    assert frames[1].parent_frame_id == frames[0].frame_id
    assert frames[1].source_coordinates["start_column"] == 2
    assert frames[0].source_version == source["version"]


def test_index_is_metadata_only_and_exports_original_evidence(spreadsheet, tmp_path):
    service, asset = indexed(spreadsheet, tmp_path)
    index = service.list_frame_index()
    assert index[0]["source_coordinates"]["type"] == spreadsheet.suffix[1:]
    assert "content" not in index[0] and "original_content" not in index[0]
    frame = service._resolve_frame(ContextFrame.from_index(index[0]))
    exported = ContextPackage("context_ready", "values", [frame]).to_dict()
    assert exported["frames"][0]["content"] == frame.content
    assert exported["frames"][0]["source_version"] == index[0]["source_version"]
    assert frame.content.startswith("A1 = ")
    # JSON persistence/reload keeps coordinates and source version.
    reloaded = ContextSubagent(service.workspace)
    assert reloaded.list_frame_index() == index


def test_changed_original_file_rejects_even_when_cell_values_unchanged(spreadsheet, tmp_path):
    service, _ = indexed(spreadsheet, tmp_path)
    indexed_frame = ContextFrame.from_index(service.list_frame_index()[0])
    if spreadsheet.suffix == ".csv":
        # Equivalent records, different source bytes.
        spreadsheet.write_bytes(b"\xef\xbb\xbf" + spreadsheet.read_bytes())
    else:
        from openpyxl import load_workbook
        wb = load_workbook(spreadsheet)
        wb.properties.title = "changed metadata"
        wb.save(spreadsheet)
    with pytest.raises(ValueError, match="Source changed"):
        service._resolve_frame(indexed_frame)


@pytest.mark.parametrize("mutation", [
    {"start_row": 0}, {"end_row": 999}, {"start_column": True}, {"end_column": 99},
    {"sheet": "invented"}, {"column_ids": ["wrong"]}, {"type": "paragraph"},
])
def test_invalid_coordinates_fail(spreadsheet, mutation):
    source = read_spreadsheet(spreadsheet)
    c = {**grid_coordinates(source, source["grids"][0]), **mutation}
    with pytest.raises(ValueError):
        resolve_spreadsheet(source, c, source["version"])


def test_parent_child_containment_coverage_and_unknown_keys(spreadsheet):
    source = read_spreadsheet(spreadsheet)
    ranges = native_ranges(source)
    assert validate_grid_ranges(source, ranges)[1]["depth"] == 1
    for bad in (
        [{**item, "parentKey": "unknown"} for item in ranges],
        [{**item, "coordinates": ranges[0]["coordinates"]} for item in ranges],
        ranges[1:],
        ranges + [ranges[0]],
    ):
        with pytest.raises(ValueError):
            validate_grid_ranges(source, bad)


def test_selection_rejects_unknown_ids_and_model_evidence(spreadsheet, tmp_path, monkeypatch):
    service, _ = indexed(spreadsheet, tmp_path)
    candidates = [ContextFrame.from_index(entry) for entry in service.list_frame_index()]
    def choose(prompt, system):
        request_id = prompt.split("Request ID: ")[1].split("\n")[0]
        return {"request_id": request_id, "selected_frames": [
            {"context_frame_id": "invented", "include_mode": "full_frame", "priority": 1},
        ]}
    monkeypatch.setattr(service.models.adapter, "generate_json", choose)
    with pytest.raises(ValueError, match="Unknown frame"):
        service._select_frames("values", candidates, 5, "stub")


def test_parent_child_duplicate_evidence_and_same_rows_different_columns(spreadsheet, tmp_path, monkeypatch):
    source = read_spreadsheet(spreadsheet)
    service = ContextSubagent(tmp_path / "workspace")
    asset = service.register_asset(str(spreadsheet))
    frames = service.frame_builder.build_spreadsheet(asset.asset_id, source, str(spreadsheet), planner=native_ranges)
    service.frames.replace_for_asset(asset.asset_id, frames)
    candidates = [ContextFrame.from_index(entry) for entry in service.list_frame_index()]
    def choose(prompt, system):
        request_id = prompt.split("Request ID: ")[1].split("\n")[0]
        return {"request_id": request_id, "selected_frames": [
            {"context_frame_id": frame.frame_id, "include_mode": "full_frame", "priority": i + 1}
            for i, frame in enumerate(candidates)]}
    monkeypatch.setattr(service.models.adapter, "generate_json", choose)
    selected = service._select_frames("values", candidates, 10, "stub")
    assert selected[0].content
    assert selected[1].content == ""
    assert selected[1].duplicate_of_frame_id == selected[0].frame_id
    if source["type"] == "xlsx":
        assert selected[2].content  # A different sheet is never a duplicate.


def test_legacy_paragraph_frames_remain_compatible(tmp_path):
    path = tmp_path / "source.md"
    path.write_text("first\n\nsecond")
    service = ContextSubagent(tmp_path / "workspace")
    asset = service.register_asset(str(path))
    service.analyze_asset(asset.asset_id)
    entries = service.frames.store.load()
    for entry in entries:
        entry.pop("source_coordinates", None)
        entry.pop("source_version", None)
    service.frames.store.save(entries)
    service.frames._save(service.frames.list())
    frame = service._resolve_frame(ContextFrame.from_index(service.list_frame_index()[0]))
    assert frame.content == "first\n\nsecond"


def test_native_refinement_preserves_parent_coordinates_and_other_sheets(spreadsheet, tmp_path, monkeypatch):
    service, asset = indexed(spreadsheet, tmp_path)
    original_frames = service.list_frames(asset.asset_id)
    parent = original_frames[0]
    service.settings.save({"selected_model_id": "stub"})
    monkeypatch.setattr(service.models, "select", lambda model: model)
    monkeypatch.setattr(service, "_frame_summarizer", lambda model: lambda content, location: "Range metadata")
    monkeypatch.setattr(service, "_spreadsheet_planner", native_ranges)
    refined = service.refine_frame(parent.frame_id, "Useful value columns")
    assert len(refined) == 2
    root = next(frame for frame in refined if frame.coordinates() == parent.coordinates())
    assert root.source_version == parent.source_version
    assert root.revision_parent_frame_id == parent.frame_id
    assert root.version == parent.version + 1
    assert not service.get_frame(parent.frame_id).is_current
    if len(original_frames) > 1:
        assert service.get_frame(original_frames[1].frame_id).is_current


def test_cross_sheet_parent_and_same_rows_disjoint_columns_are_not_duplicates(tmp_path):
    path = tmp_path / "sheets.xlsx"
    wb = Workbook()
    wb.active.title = "One"
    wb.active.append(["left", "right"])
    wb.create_sheet("Two").append(["other", "values"])
    wb.save(path)
    source = read_spreadsheet(path)
    one = grid_coordinates(source, source["grids"][0])
    two = grid_coordinates(source, source["grids"][1])
    with pytest.raises(ValueError, match="strictly"):
        validate_grid_ranges(source, [
            {"key": "one", "parentKey": None, "coordinates": one, "reason": "Full table"},
            {"key": "two", "parentKey": "one", "coordinates": two, "reason": "Wrong sheet"},
        ])
    from langform.coordinates import contains
    left = grid_coordinates(source, source["grids"][0], 1, 1, 1, 1)
    right = grid_coordinates(source, source["grids"][0], 1, 1, 2, 2)
    assert not contains(left, right)
    assert not contains(right, left)
