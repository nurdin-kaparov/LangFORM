"""Windows executable entry point. Python and dependencies are bundled at build time."""
from __future__ import annotations

import argparse
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile

import launcher
from langform.bridge import main as bridge_main


def configuration_path() -> Path:
    """Configuration sits beside the real executable, not in PyInstaller's temp dir."""
    executable = Path(sys.executable if getattr(sys, "frozen", False) else __file__)
    return executable.resolve().parent / "langform-connection.json"


def resolve_server(explicit: str | None, config: Path) -> str:
    if explicit:
        return launcher.validate_server(explicit)
    if config.is_file():
        data = json.loads(config.read_text(encoding="utf-8"))
        if not isinstance(data, dict) or not isinstance(data.get("server"), str):
            raise ValueError("langform-connection.json must contain a server address.")
        return launcher.validate_server(data["server"])
    return launcher.validate_server(input("LangFORM website address (https://...): ").strip())


def read_settings(path: Path) -> dict:
    if not path.is_file():
        return {}
    data = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(data, dict):
        raise ValueError("Saved connection settings are invalid.")
    return data


def smoke_test() -> int:
    """Offline check used by Windows CI, with no pairing, network, or folder grants."""
    import ssl
    import tkinter
    import requests
    from docx import Document
    from openpyxl import Workbook, load_workbook
    from pypdf import PdfReader, PdfWriter
    from langform.bridge import capability_report
    from langform.bridge_local import LocalBridge

    with tempfile.TemporaryDirectory(prefix="langform-smoke-") as directory:
        root = Path(directory)
        book = Workbook()
        book.active["A1"] = "LangFORM"
        book.save(root / "sample.xlsx")
        assert load_workbook(root / "sample.xlsx").active["A1"].value == "LangFORM"
        document = Document()
        document.add_paragraph("LangFORM")
        document.save(root / "sample.docx")
        assert Document(root / "sample.docx").paragraphs[0].text == "LangFORM"
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        writer.write(root / "sample.pdf")
        assert len(PdfReader(root / "sample.pdf").pages) == 1
        assert LocalBridge(root / "workspace", [], "http://127.0.0.1:11434").sources() == {}
        assert capability_report()["supportedObjectives"] == ["processing", "packaging"]
        assert Path(requests.certs.where()).is_file()
        ssl.create_default_context()
        tcl = tkinter.Tcl()
        assert tcl.eval("info patchlevel")
    print("LangFORM executable smoke check passed: runtime, TLS, dialogs, document readers, bridge.")
    return 0


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description="Connect your Windows PC to LangFORM. No Python installation required.")
    parser.add_argument("--server", help="LangFORM HTTPS website address; otherwise read adjacent configuration or ask")
    parser.add_argument("--change-folder", action="store_true", help="Choose or remove the approved folder")
    parser.add_argument("--re-pair", action="store_true", help="Request new browser-approved pairing")
    parser.add_argument("--smoke-test", action="store_true", help="Offline packaging check; never connects or pairs")
    args = parser.parse_args(argv)
    if args.smoke_test:
        return smoke_test()
    server = resolve_server(args.server, configuration_path())
    home = Path.home() / ".langform"
    settings_path = home / "windows-companion-settings.json"
    settings = read_settings(settings_path)
    same_server = settings.get("server") == server
    if not same_server and not launcher.ask(
        f"Connect this PC to {server}?\n\n"
        "This executable includes Python and the LangFORM bridge. No Python or package installation "
        "is needed. Existing Ollama and models will be reused; no model is downloaded. "
        "Folder access is optional and browser pairing still requires your approval."
    ):
        print("Connection cancelled. No pairing or folder access was performed.")
        return 0
    if same_server and not args.change_folder:
        folder = settings.get("folder")
        if folder is not None and not isinstance(folder, str):
            raise ValueError("Saved folder settings are invalid. Use --change-folder to choose a new grant.")
        if folder and not Path(folder).is_dir():
            print("The approved folder is unavailable. Connecting without file access.")
            folder = None
    else:
        folder = launcher.choose_folder()
    if not launcher.ollama_reachable():
        command = launcher.existing_ollama()
        if command and launcher.ask("Ollama is installed but not responding. Start your existing Ollama server?"):
            subprocess.Popen([command, "serve"], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        else:
            print("Open or install Ollama on this PC. LangFORM does not install it or download models automatically.")
    home.mkdir(parents=True, exist_ok=True, mode=0o700)
    # Settings contain only the approved server and optional folder, never a device credential.
    settings_path.write_text(json.dumps({"server": server, "folder": folder}), encoding="utf-8")
    os.chmod(settings_path, 0o600)
    arguments = ["--server", server, "--workspace", str(home / "companion-workspace"),
                 "--open-browser", "--label", "Windows Context SLM"]
    if folder:
        arguments += ["--allow", folder]
    if args.re_pair:
        arguments.append("--re-pair")
    print("Starting LangFORM. Keep this window open; Ctrl+C disconnects.")
    return bridge_main(arguments)


if __name__ == "__main__":
    code = 1
    try:
        code = main()
    except KeyboardInterrupt:
        print("LangFORM stopped.")
        code = 0
    except Exception as error:
        print(f"LangFORM could not start: {error}")
    finally:
        # Keep double-clicked console failures and pairing messages visible.
        if getattr(sys, "frozen", False) and "--smoke-test" not in sys.argv:
            try:
                input("\nPress Enter to close...")
            except (EOFError, KeyboardInterrupt):
                pass
    raise SystemExit(code)
